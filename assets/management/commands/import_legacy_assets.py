"""
One-time migration: import the old multi-sheet 'Assets.xlsx' tracking file
into the Asset Register database.

This is NOT the same as the admin "Bulk Import" feature in the web UI (that
one expects a single clean sheet in the app's own format). This command is
written specifically for the legacy file's messy, inconsistent layout —
each old sheet has its own column shape and needs its own parsing rules.

Usage:
    python manage.py import_legacy_assets /path/to/Assets.xlsx --dry-run
    python manage.py import_legacy_assets /path/to/Assets.xlsx

Always run with --dry-run first and read the report before committing.

Sheets handled: keyboard, Monitor, Mouse, UPS, Hard disk, Bluetooth,
Software and OS, others, WorkStation_List.

Sheets intentionally SKIPPED (not per-asset records, need manual review):
System, Project Details, Project backup, inside the Cupboard,
IT_Vendor List, Incident Register.
"""
import re
from datetime import datetime

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from assets.models import Asset, AssetCategory
from assets.import_utils import STATUS_LOOKUP, _normalize_header

SKIPPED_SHEETS = {
    "System": "no headers / no stable IDs — 4 cryptic rows, review manually",
    "Project Details": "project list, not asset records",
    "Project backup": "project backup log, not asset records",
    "inside the Cupboard": "spare/unlabeled internal HDDs with no stable ID",
    "IT_Vendor List": "vendor contact list, not asset records",
    "Incident Register": "incident log, not asset records",
}

DATE_PATTERNS = [
    r"(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})",
]
DATE_FORMATS = ["%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y"]


def clean(v):
    if v is None:
        return ""
    return str(v).strip()


def guess_status(text, default="working"):
    text = clean(text).lower()
    if not text:
        return default
    mapped = STATUS_LOOKUP.get(_normalize_header(text))
    if mapped:
        return mapped
    if "not working" in text:
        return "not_working"
    if "scrap" in text or "destroy" in text:
        return "scrap"
    if "missing" in text:
        return "missing"
    if "idle" in text or "cupboard" in text:
        return "idle"
    if "service" in text:
        return "service"
    if "working" in text:
        return "working"
    return "other"


def guess_assignment(text):
    """Best-effort split of a free-text remark into (assigned_to, location, leftover_note).
    Defaults to putting the whole thing in the leftover note when unsure, rather
    than guessing wrong."""
    text = clean(text)
    if not text:
        return "", "", ""
    low = text.lower()
    if low.startswith("used by"):
        return text[7:].strip(" :-"), "", ""
    if low.startswith("used in"):
        return "", text[7:].strip(" :-"), ""
    if "cupboard" in low:
        return "", "In Cupboard", "" if low == "in cupboard" else text
    return "", "", text


def find_date(text):
    text = clean(text)
    for pattern in DATE_PATTERNS:
        m = re.search(pattern, text)
        if not m:
            continue
        candidate = m.group(1)
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(candidate, fmt).date()
            except ValueError:
                continue
    return None


class LegacyImporter:
    def __init__(self, wb, dry_run=True, stdout=None):
        self.wb = wb
        self.dry_run = dry_run
        self.stdout = stdout
        self.categories = {c.name.strip().lower(): c for c in AssetCategory.objects.all()}
        self.existing_tags = set(Asset.objects.values_list("asset_tag", flat=True))
        self.created = 0
        self.skipped_rows = []   # (sheet, row_num, reason)
        self.created_rows = []   # (sheet, tag, name)

    def log(self, msg):
        if self.stdout:
            self.stdout.write(msg)

    def get_category(self, *names):
        for name in names:
            cat = self.categories.get(name.strip().lower())
            if cat:
                return cat
        return None

    def make_asset(self, sheet, row_num, category, tag, name, **kwargs):
        tag = clean(tag)
        if not category:
            self.skipped_rows.append((sheet, row_num, f"no matching category for tag '{tag}'"))
            return
        if not tag:
            self.skipped_rows.append((sheet, row_num, "missing asset tag"))
            return
        if tag in self.existing_tags:
            self.skipped_rows.append((sheet, row_num, f"asset tag '{tag}' already exists — skipped"))
            return
        self.existing_tags.add(tag)

        fields = dict(
            category_id=category.id,
            asset_tag=tag[:50],
            name=(name or tag)[:150],
            brand=clean(kwargs.get("brand"))[:100],
            model_number=clean(kwargs.get("model_number"))[:100],
            serial_number=clean(kwargs.get("serial_number"))[:150],
            status=kwargs.get("status") or "working",
            current_assigned_to=clean(kwargs.get("current_assigned_to"))[:150],
            current_location=clean(kwargs.get("current_location"))[:150],
            purchase_date=kwargs.get("purchase_date"),
            last_service_date=kwargs.get("last_service_date"),
            notes=clean(kwargs.get("notes"))[:2000],
        )
        self.created_rows.append((sheet, fields["asset_tag"], fields["name"]))
        self.created += 1
        if not self.dry_run:
            Asset.objects.create(**fields)

    def sheet_rows(self, name):
        if name not in self.wb.sheetnames:
            return []
        rows = []
        for r in self.wb[name].iter_rows(values_only=True):
            if any(c is not None and str(c).strip() for c in r):
                rows.append(r)
        return rows

    # ---- per-sheet handlers -------------------------------------------------

    def import_keyboard(self):
        cat = self.get_category("Keyboard")
        for i, row in enumerate(self.sheet_rows("keyboard")[1:], start=2):
            row = list(row) + [None] * 6
            tag, brand, status = row[0], row[1], row[2]
            remark = row[5]
            assigned, loc, note = guess_assignment(remark)
            self.make_asset(
                "keyboard", i, cat, tag, f"{clean(brand)} Keyboard".strip(),
                brand=brand, status=guess_status(status),
                current_assigned_to=assigned, current_location=loc, notes=note,
            )

    def import_monitor(self):
        cat = self.get_category("Monitor")
        for i, row in enumerate(self.sheet_rows("Monitor")[1:], start=2):
            row = list(row) + [None] * 9
            tag, brand, m_type, color, size, model, serial, condition, extra = row[:9]
            assigned, loc, note = guess_assignment(extra)
            desc = ", ".join(x for x in [clean(m_type), clean(color), clean(size)] if x)
            self.make_asset(
                "Monitor", i, cat, tag, f"{clean(brand)} Monitor".strip(),
                brand=brand, model_number=model, serial_number=serial,
                status=guess_status(condition),
                current_assigned_to=assigned, current_location=loc,
                notes=", ".join(x for x in [desc, note] if x),
            )

    def import_mouse(self):
        cat = self.get_category("Mouse")
        for i, row in enumerate(self.sheet_rows("Mouse")[1:], start=2):
            row = list(row) + [None] * 6
            tag, brand, status, extra_note, serial, remark = row[:6]
            assigned, loc, note = guess_assignment(remark)
            purchase_date = find_date(extra_note)
            self.make_asset(
                "Mouse", i, cat, tag, f"{clean(brand)} Mouse".strip(),
                brand=brand, serial_number=serial, status=guess_status(status),
                current_assigned_to=assigned, current_location=loc,
                purchase_date=purchase_date,
                notes=", ".join(x for x in [clean(extra_note), note] if x),
            )

    def import_ups(self):
        cat = self.get_category("UPS")
        for i, row in enumerate(self.sheet_rows("UPS")[1:], start=2):
            row = list(row) + [None] * 9
            tag, brand, color, model, size, last_service, condition, status, current = row[:9]
            assigned, loc, note = guess_assignment(current)
            self.make_asset(
                "UPS", i, cat, tag, f"{clean(brand)} UPS".strip(),
                brand=brand, model_number=model,
                status=guess_status(condition or status),
                current_assigned_to=assigned, current_location=loc,
                last_service_date=find_date(last_service),
                notes=", ".join(x for x in [clean(color), clean(size), clean(last_service), note] if x),
            )

    def import_hard_disk(self):
        cat = self.get_category("Hard Disk")
        for i, row in enumerate(self.sheet_rows("Hard disk")[1:], start=2):
            row = list(row) + [None] * 6
            size, serial, conditions, purpose, status, current = row[:6]
            if not clean(serial):
                self.skipped_rows.append(("Hard disk", i, "no serial number to use as asset tag"))
                continue
            assigned, loc, note = guess_assignment(current)
            self.make_asset(
                "Hard disk", i, cat, serial, f"{clean(size)} Hard Disk".strip(),
                serial_number=serial, status=guess_status(status or conditions),
                current_assigned_to=assigned, current_location=loc,
                notes=", ".join(x for x in [clean(conditions), clean(purpose), note] if x),
            )

    def import_bluetooth(self):
        cat = self.get_category("Bluetooth Device", "Other Asset", "Other Assets")
        for i, row in enumerate(self.sheet_rows("Bluetooth")[1:], start=2):
            row = list(row) + [None] * 5
            tag, devices, brand, remark, user = row[:5]
            self.make_asset(
                "Bluetooth", i, cat, tag, clean(devices) or "Bluetooth Device",
                brand=brand, current_assigned_to=user, notes=remark,
            )

    def import_software(self):
        cat = self.get_category("Software / OS License", "Software/OS License")
        for i, row in enumerate(self.sheet_rows("Software and OS")[1:], start=2):
            row = list(row) + [None] * 5
            s_type, version, product_key, system_no, extra = row[:5]
            name = f"{clean(s_type)} {clean(version)}".strip()
            note = ", ".join(x for x in [f"System: {clean(system_no)}" if system_no else "", clean(extra)] if x)
            self.make_asset(
                "Software and OS", i, cat, product_key, name,
                serial_number=product_key, notes=note,
            )

    OTHER_DEVICE_CATEGORY = {
        "a/c": "Air Conditioner",
    }

    def import_others(self):
        for i, row in enumerate(self.sheet_rows("others")[1:], start=2):
            row = list(row) + [None] * 5
            tag, device_type, device_name, spec, note = row[:5]
            cat_name = self.OTHER_DEVICE_CATEGORY.get(clean(device_type).lower())
            cat = self.get_category(cat_name) if cat_name else None
            if not cat:
                cat = self.get_category("Other Asset", "Other Assets")
            self.make_asset(
                "others", i, cat, tag, f"{clean(device_type)} - {clean(device_name)}".strip(" -"),
                brand=device_name, current_location=spec,
                last_service_date=find_date(note), notes=note,
            )

    def import_workstations(self):
        cat = self.get_category("Workstation")
        header = None
        records = {}  # tag -> (row_num, data dict) — later occurrences overwrite earlier ones
        for i, row in enumerate(self.sheet_rows("WorkStation_List"), start=1):
            row = list(row)
            if clean(row[0]) == "Workstation ID":
                header = [clean(h) for h in row]
                continue
            if not header:
                continue
            data = dict(zip(header, row))
            tag = clean(data.get("Workstation ID"))
            if not tag:
                continue
            records[tag] = (i, data)

        for tag, (i, data) in records.items():
            employee_id = clean(data.get("Employee ID"))
            employee_name = clean(data.get("Employee Name"))
            assigned_to = employee_name or employee_id
            if employee_id.lower() in ("no one", "none", ""):
                assigned_to = ""
                status = "idle"
            else:
                status = "working"
            components = []
            for label in ["CPU Number", "Monitor Number", "Keyboard Number", "Mouse Number", "UPS No."]:
                val = clean(data.get(label))
                if val and val.lower() != "none":
                    components.append(f"{label}: {val}")
            for label in ["Product Id", "Cd Key", "Operating System", "Purposes", "User Id"]:
                val = clean(data.get(label))
                if val:
                    components.append(f"{label}: {val}")
            self.make_asset(
                "WorkStation_List", i, cat, tag,
                f"Workstation - {employee_name}" if employee_name else "Workstation",
                status=status, current_assigned_to=assigned_to,
                notes="; ".join(components),
            )

    def run(self):
        self.import_keyboard()
        self.import_monitor()
        self.import_mouse()
        self.import_ups()
        self.import_hard_disk()
        self.import_bluetooth()
        self.import_software()
        self.import_others()
        self.import_workstations()


class Command(BaseCommand):
    help = "One-time import of the legacy multi-sheet Assets.xlsx into the Asset Register."

    def add_arguments(self, parser):
        parser.add_argument("file_path", type=str, help="Path to the legacy Assets.xlsx file")
        parser.add_argument("--dry-run", action="store_true", help="Preview without saving anything")

    def handle(self, *args, **options):
        path = options["file_path"]
        dry_run = options["dry_run"]

        try:
            wb = load_workbook(filename=path, data_only=True)
        except Exception as exc:
            raise CommandError(f"Couldn't open '{path}': {exc}")

        importer = LegacyImporter(wb, dry_run=dry_run, stdout=self.stdout)

        if dry_run:
            importer.run()
        else:
            with transaction.atomic():
                importer.run()

        self.stdout.write("")
        self.stdout.write(self.style.SUCCESS(
            f"{'[DRY RUN] Would create' if dry_run else 'Created'} {importer.created} asset(s)."
        ))

        if importer.created_rows:
            self.stdout.write("\nSample of created assets (first 20):")
            for sheet, tag, name in importer.created_rows[:20]:
                self.stdout.write(f"  [{sheet}] {tag} — {name}")

        if importer.skipped_rows:
            self.stdout.write(self.style.WARNING(f"\n{len(importer.skipped_rows)} row(s) skipped:"))
            for sheet, row_num, reason in importer.skipped_rows[:50]:
                self.stdout.write(f"  [{sheet} row {row_num}] {reason}")
            if len(importer.skipped_rows) > 50:
                self.stdout.write(f"  ... and {len(importer.skipped_rows) - 50} more")

        self.stdout.write(self.style.WARNING(
            "\nSheets NOT processed (not per-asset records — review/enter manually):"
        ))
        for sheet, reason in SKIPPED_SHEETS.items():
            if sheet in wb.sheetnames:
                self.stdout.write(f"  [{sheet}] {reason}")

        if dry_run:
            self.stdout.write(self.style.NOTICE(
                "\nThis was a dry run — nothing was saved. Re-run without --dry-run to commit."
            ))
